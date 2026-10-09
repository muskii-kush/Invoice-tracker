"""
Read/write layer for Invoice Automation Tracker.xlsx. Every write is an
append (or a targeted single-cell status update) — nothing here ever
overwrites a Remarks/Activity cell, per the master prompt's
"do not overwrite previous remarks, add a new activity entry instead" rule.
"""
from openpyxl import load_workbook

from . import schema


def open_tracker(path):
    return load_workbook(path)


def save_tracker(wb, path):
    wb.save(path)


def _header_index(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def next_id(ws, prefix):
    """Sequential ID like INV-0001, based on the current row count."""
    n = ws.max_row  # header is row 1, so this is (existing records + 1)
    return f"{prefix}-{n:04d}"


def append_row(wb, sheet_name, data: dict):
    """
    data: dict keyed by column header (must match schema.SHEETS[sheet_name]
    exactly — unknown keys are ignored, missing keys are left blank).
    Returns the generated ID (for sheets with an ID column) or None.
    """
    ws = wb[sheet_name]
    headers = schema.SHEETS[sheet_name]
    id_col = headers[0]
    prefix = schema.ID_PREFIXES.get(sheet_name)

    row_id = None
    if prefix and not data.get(id_col):
        row_id = next_id(ws, prefix)
        data = {**data, id_col: row_id}
    elif data.get(id_col):
        row_id = data[id_col]

    row = [data.get(h, "") for h in headers]
    ws.append(row)
    return row_id


def find_row_by_id(ws, id_column_index, target_id):
    for row in ws.iter_rows(min_row=2):
        if row[id_column_index - 1].value == target_id:
            return row
    return None


def update_cell(wb, sheet_name, row_id, column_name, value):
    """Targeted update of one field (e.g. Intake Status, P2P Status,
    Matched Invoice ID) — never used for Remarks/free-text history."""
    ws = wb[sheet_name]
    headers = schema.SHEETS[sheet_name]
    id_col_idx = 1
    col_idx = headers.index(column_name) + 1
    row = find_row_by_id(ws, id_col_idx, row_id)
    if row is None:
        raise ValueError(f"{row_id} not found in {sheet_name}")
    row[col_idx - 1].value = value


def all_invoices(wb):
    """Return every Invoice Register row as a list of dicts, for matcher.match()."""
    ws = wb["Invoice Register"]
    headers = schema.INVOICE_REGISTER
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            continue
        record = dict(zip(headers, row))
        out.append({
            "invoice_id": record["Invoice ID"],
            "vendor": record["Vendor"],
            "invoice_number": record["Invoice Number"],
            "amount": record["Invoice Amount"],
            "invoice_date": record["Invoice Date"],
            "billing_period": record["Billing Period"],
            "po_number": record["PO Number"],
            "pr_number": record["PR Number"],
            "p2p_link": record["P2P / Aerchain Link"],
            "attachment_hash": None,  # populated by callers that hash attachments
            "description": record["Description"],
        })
    return out


def known_vendors(wb):
    ws = wb["Vendor Reference"]
    headers = schema.VENDOR_REFERENCE
    vendors = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in (None, ""):
            continue
        record = dict(zip(headers, row))
        if str(record.get("Active / Inactive", "Active")).lower() != "inactive":
            vendors.append(record["Vendor"])
    return vendors
